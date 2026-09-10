"""Anonymous regressions extracted from the private corpus audit."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document as WordDocument
from pydantic import ValidationError

from document_parser import (
    ContainerBlock,
    HeadingBlock,
    MarkdownOptions,
    MarkdownProfile,
    ParagraphBlock,
    TableBlock,
    TableCell,
    TableMode,
    TableRow,
    TextSpan,
    convert,
    normalize_document,
    to_markdown,
)
from document_parser.pdf_adapter import _native_quality, _PdfWord
from test_phase3_helpers import make_document

FIXTURES = Path(__file__).parent / "fixtures" / "synthetic"


@pytest.mark.parametrize("split", range(1, 30))
def test_run_boundary_never_changes_plain_text(split: int) -> None:
    text = "Send/Receive menyusuna keçin. Work Offline"
    block = ParagraphBlock(
        block_id="p", spans=(TextSpan(text=text[:split], bold=True), TextSpan(text=text[split:]))
    )
    normalized = normalize_document(make_document(block))
    paragraph = normalized.blocks[0]
    assert isinstance(paragraph, ParagraphBlock)
    assert "".join(span.text for span in paragraph.spans) == text
    assert normalize_document(normalized).model_dump_json() == normalized.model_dump_json()


def test_clean_profile_is_default_and_rejects_html_options() -> None:
    assert MarkdownOptions().profile is MarkdownProfile.CLEAN
    for options in ({"table_mode": "html"}, {"include_source_markers": True}):
        with pytest.raises(ValidationError, match="fidelity"):
            MarkdownOptions(**options)
    assert MarkdownOptions(profile=MarkdownProfile.FIDELITY, table_mode=TableMode.HTML)


@pytest.mark.parametrize(
    "name", ("rich-structure.docx", "digital-native.pdf", "structured-workbook.xlsx")
)
def test_clean_conversion_has_no_generated_html_or_asset_links(name: str) -> None:
    result = convert(FIXTURES / name)
    for marker in ("<table", "<u>", "<br>", "<!--", "!["):
        assert marker not in result.markdown
    assert result.document.quality is not None
    assert result.document.quality.overall_score < 1
    assert convert(FIXTURES / name).markdown == result.markdown


def test_clean_preserves_literals_and_contextual_punctuation() -> None:
    document = make_document(
        ParagraphBlock(
            block_id="p",
            spans=(TextSpan(text="IP 10.0.0.1 - Work Offline <field>", underline=True),),
        )
    )
    assert to_markdown(document) == "IP 10.0.0.1 - Work Offline \\<field\\>\n"
    assert "<u>" in to_markdown(document, options=MarkdownOptions(profile=MarkdownProfile.FIDELITY))


def test_record_table_keeps_merged_context_and_nested_content() -> None:
    table = TableBlock(
        block_id="t",
        row_count=3,
        column_count=2,
        rows=(
            TableRow(
                row_index=0,
                cells=(
                    TableCell(column_index=0, is_header=True, displayed_text="Group"),
                    TableCell(column_index=1, is_header=True, displayed_text="Detail"),
                ),
            ),
            TableRow(
                row_index=1,
                cells=(
                    TableCell(column_index=0, row_span=2, displayed_text="North"),
                    TableCell(column_index=1, displayed_text="Alpha"),
                ),
            ),
            TableRow(
                row_index=2,
                cells=(
                    TableCell(
                        column_index=1,
                        blocks=(ParagraphBlock(block_id="nested", spans=(TextSpan(text="Beta"),)),),
                    ),
                ),
            ),
        ),
    )
    markdown = to_markdown(make_document(table))
    assert markdown.count("North") == 2
    assert "Alpha" in markdown and "Beta" in markdown
    assert "**Group:**" in markdown and "**Detail:**" in markdown
    assert "<" not in markdown


def test_numbered_docx_heading_is_not_a_list_item() -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    word = WordDocument()
    heading = word.add_paragraph("Purpose", style="Heading 1")
    num = OxmlElement("w:numPr")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), "5")
    num.append(number_id)
    heading._p.get_or_add_pPr().append(num)
    word.add_paragraph("Plain styled paragraph", style="List Paragraph")
    stream = io.BytesIO()
    word.save(stream)
    result = convert(stream.getvalue(), filename="numbered.docx")
    section = result.document.blocks[0]
    assert isinstance(section, ContainerBlock)
    assert isinstance(section.blocks[0], HeadingBlock)
    assert isinstance(section.blocks[1], ParagraphBlock)
    assert "\n    #" not in result.markdown


def test_native_text_quality_distinguishes_mixed_script_tokens() -> None:
    def word(text: str) -> _PdfWord:
        return _PdfWord(text, 0, 100, 0, 12, 12, "Font")

    assert _native_quality((word("English"), word("Русский")))[0] == 0.95
    assert _native_quality((word("Camлyat"), word("normal")))[0] < 0.7
    assert _native_quality(())[0] == 0


def test_reconciliation_handles_one_native_paragraph_and_multiple_ocr_lines() -> None:
    from test_normalization_quality import document_with, page, paragraph, point_box

    native = paragraph(
        "native", "Alpha Beta", method="native", box=point_box().model_copy(update={"height": 30})
    )
    first = paragraph("first", "Alpha", method="ocr", confidence=0.9, box=point_box())
    second = paragraph("second", "Beta", method="ocr", confidence=0.9, box=point_box(y=15))
    result = normalize_document(document_with(page(first, second, native)))
    container = result.blocks[0]
    assert isinstance(container, ContainerBlock)
    assert [b.attributes["active_for_rag"] for b in container.blocks] == [True, True, False]
    assert result.diagnostics == ()
    assert normalize_document(result).model_dump_json() == result.model_dump_json()


def test_cell_center_assignment_does_not_leak_a_neighboring_glyph() -> None:
    from collections.abc import Callable
    from types import SimpleNamespace
    from typing import Any

    from document_parser.pdf_adapter import _materialize_table

    class Page:
        def __init__(self, objects: list[dict[str, Any]]) -> None:
            self.objects = objects

        def filter(self, predicate: Callable[[dict[str, Any]], bool]) -> Page:
            return Page([obj for obj in self.objects if predicate(obj)])

        def extract_text(self) -> str:
            return "".join(obj["text"] for obj in self.objects)

    page = Page(
        [
            {"object_type": "char", "x0": 1, "x1": 5, "top": 7, "bottom": 11, "text": "A"},
            {"object_type": "char", "x0": 1, "x1": 5, "top": 11, "bottom": 19, "text": "B"},
            {"object_type": "rect", "text": "must not extract"},
        ]
    )
    table = SimpleNamespace(cells=[(0, 0, 20, 10), (0, 10, 20, 20)], bbox=(0, 0, 20, 20))
    assert [cell.text for cell in _materialize_table(table, page).cells] == ["A", "B"]
