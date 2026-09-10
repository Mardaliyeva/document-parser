"""Anonymous boundary tests for native extraction and clean serialization."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

import pytest
from docx import Document as WordDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from document_parser import (
    ContainerBlock,
    ContainerRole,
    ContentBlock,
    HeadingBlock,
    ListBlock,
    ListItem,
    ListKind,
    MarkdownOptions,
    OcrConfigurationError,
    OcrDependencyNotAvailableError,
    OcrOptions,
    ParagraphBlock,
    ParseOptions,
    SourceLocation,
    TableBlock,
    TableCell,
    TableMode,
    TableRow,
    TextSpan,
    normalize_document,
    to_markdown,
)
from document_parser.docx_adapter import (
    _DocxContext,
    _list_spec,
    _number_label,
    _number_patterns,
    _numbering_map,
)
from document_parser.markdown import _render_span
from document_parser.normalization import _confidence
from document_parser.pdf_adapter import (
    _mark_repeated_tables,
    _paragraphs_and_lists,
    _pdf_table,
    _PdfCellData,
    _PdfContext,
    _PdfPageData,
    _PdfTableData,
)
from test_normalization_quality import document_with, page, paragraph, point_box
from test_paddle_ocr import FakePaddleModule, install_modules
from test_phase3_helpers import make_document


@pytest.mark.parametrize(
    ("value", "fmt", "expected"),
    [
        (27, "lowerLetter", "aa"),
        (52, "upperLetter", "AZ"),
        (1949, "lowerRoman", "mcmxlix"),
        (4, "upperRoman", "IV"),
        (10000, "upperRoman", "10000"),
        (7, "decimal", "7"),
    ],
)
def test_number_labels(value: int, fmt: str, expected: str) -> None:
    assert _number_label(value, fmt) == expected


def test_number_overrides_restarts_and_explicit_no_numbering() -> None:
    word = WordDocument()
    root = word.part.numbering_part.element
    root.append(
        parse_xml(f"""<w:abstractNum {nsdecls("w")} w:abstractNumId="80">
      <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>
      <w:lvlText w:val="%1."/></w:lvl><w:lvl w:ilvl="1"/>
    </w:abstractNum>""")
    )
    root.append(
        parse_xml(f"""<w:num {nsdecls("w")} w:numId="80">
      <w:abstractNumId w:val="80"/>
      <w:lvlOverride w:ilvl="0"><w:startOverride w:val="3"/>
        <w:lvl w:ilvl="0"><w:start w:val="2"/><w:numFmt w:val="upperRoman"/>
        <w:lvlText w:val="%1."/><w:lvlRestart w:val="0"/></w:lvl>
      </w:lvlOverride>
      <w:lvlOverride w:ilvl="1"><w:lvl w:ilvl="1"/></w:lvlOverride>
      <w:lvlOverride w:ilvl="2"><w:lvl w:ilvl="2"><w:numFmt w:val="bullet"/>
      </w:lvl></w:lvlOverride>
      <w:lvlOverride w:ilvl="3"><w:startOverride w:val="5"/></w:lvlOverride>
    </w:num>""")
    )
    root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="81"/>'))
    context = _DocxContext(ParseOptions(), "sample.docx", _numbering_map(word))
    context.number_patterns = _number_patterns(word)
    assert context.numbering[(80, 0)] == (ListKind.ORDERED, 3)
    assert context.numbering[(80, 2)][0] is ListKind.UNORDERED
    context.counters[(80, 1)] = 10
    item = word.add_paragraph("Restart")
    properties = item._p.get_or_add_pPr().get_or_add_numPr()
    properties.get_or_add_numId().val = 80
    spec = _list_spec(item, context)
    assert spec and spec.label == "III."
    assert (80, 1) not in context.counters
    properties.numId.val = 0
    assert _list_spec(item, context) is None
    # Style names are only fallback evidence for actual bullet/number styles.
    style = word.styles.add_style("List Number 9", 1)
    item.style = style
    item._p.get_or_add_pPr().remove(properties)
    assert _list_spec(item, context) is not None

    from document_parser.docx_adapter import _story_blocks

    heading = word.add_paragraph("Unnumbered bullet heading", style="Heading 1")
    number = heading._p.get_or_add_pPr().get_or_add_numPr()
    number.get_or_add_numId().val = 80
    number.get_or_add_ilvl().val = 2
    assert isinstance(_story_blocks((heading,), context)[0], HeadingBlock)


def test_docx_without_numbering_part() -> None:
    class Part:
        @property
        def numbering_part(self) -> None:
            raise KeyError("no numbering part")

    assert _number_patterns(cast(Any, SimpleNamespace(part=Part()))) == {}


def test_docx_content_controls_and_cached_fields_are_not_silently_lost() -> None:
    import io

    from document_parser import convert

    word = WordDocument()
    word.element.body.insert(
        0,
        parse_xml(f"""<w:sdt {nsdecls("w")}><w:sdtPr/>
      <w:sdtContent><w:p><w:r><w:t>Contents entry</w:t></w:r>
      <w:r><w:tab/><w:t>7</w:t></w:r></w:p></w:sdtContent></w:sdt>"""),
    )
    item = word.add_paragraph()
    item._p.append(
        parse_xml(f"""<w:fldSimple {nsdecls("w")} w:instr="DATE">
      <w:r><w:t>2026-09-08</w:t></w:r></w:fldSimple>""")
    )
    word.element.body.insert(
        1,
        parse_xml(f"""<w:del {nsdecls("w")}><w:p>
      <w:r><w:t>Deleted secret</w:t></w:r></w:p></w:del>"""),
    )
    data = io.BytesIO()
    word.save(data)
    result = convert(data.getvalue(), filename="controls.docx")
    assert "Contents entry 7" in result.markdown
    assert "2026-09-08" in result.markdown
    assert "Deleted secret" not in result.markdown


def test_clean_empty_code_links_and_numbering_markers() -> None:
    assert _render_span(TextSpan(text="  ", bold=True)) == "  "
    assert _render_span(TextSpan(text="`x`", code=True)) == "`` `x` ``"
    for url in ("javascript:alert(1)", "https://[invalid"):
        assert _render_span(TextSpan(text="link", href=url)) == "link"
    assert "%28%29" in _render_span(TextSpan(text="link", href="https://example.org/()"))
    blocks = tuple(
        ListBlock(
            block_id=f"l{i}",
            kind=ListKind.ORDERED,
            items=(
                ListItem(
                    blocks=(
                        ParagraphBlock(
                            block_id=f"p{i}",
                            spans=(TextSpan(text="First\ncontinued"),),
                            attributes={"numbering_label": label},
                        ),
                    )
                ),
            ),
        )
        for i, label in enumerate(("3)", "IV.", "1.2."))
    )
    output = to_markdown(make_document(*blocks))
    assert "3) First\n    continued" in output
    assert "- IV. First" in output
    assert "- 1.2. First" in output


def test_record_forms_multilevel_headers_empty_slots_and_fallback() -> None:
    table = TableBlock(
        block_id="table",
        row_count=3,
        column_count=3,
        rows=(
            TableRow(
                row_index=0,
                cells=(
                    TableCell(
                        column_index=0, column_span=2, is_header=True, displayed_text="Group"
                    ),
                ),
            ),
            TableRow(
                row_index=1,
                cells=(
                    TableCell(column_index=0, is_header=True, displayed_text="Group"),
                    TableCell(column_index=1, is_header=True, displayed_text="Value\nUnit"),
                ),
            ),
            TableRow(
                row_index=2,
                cells=(
                    TableCell(
                        column_index=0,
                        displayed_text="Kept",
                        blocks=(
                            ParagraphBlock(block_id="hidden", attributes={"active_for_rag": False}),
                        ),
                    ),
                    TableCell(column_index=1, displayed_text="10"),
                ),
            ),
        ),
    )
    result = to_markdown(make_document(table))
    assert "Group / Value / Unit" in result and "Kept" in result and "Sütun 3" in result
    form = TableBlock(
        block_id="form",
        row_count=1,
        column_count=2,
        rows=(
            TableRow(
                row_index=0,
                cells=(
                    TableCell(
                        column_index=0, column_span=2, displayed_text="Signature", is_header=True
                    ),
                ),
            ),
        ),
    )
    assert "Signature" in to_markdown(
        make_document(form), options=MarkdownOptions(table_mode=TableMode.RECORDS)
    )
    from document_parser import MarkdownProfile

    assert 'colspan="2"' in to_markdown(
        make_document(form), options=MarkdownOptions(profile=MarkdownProfile.FIDELITY)
    )


def test_reconcile_cycle_and_damaged_unmatched_text() -> None:
    native = tuple(paragraph(f"n{i}", "Alpha", method="native", box=point_box()) for i in range(2))
    ocr = tuple(
        paragraph(f"o{i}", "Alpha", method="ocr", confidence=0.9, box=point_box()) for i in range(2)
    )
    damaged = paragraph("bad", "garbled", method="native", box=point_box(y=60)).model_copy(
        update={"attributes": {"extraction_method": "native", "native_extraction_quality": 0.3}}
    )
    assert _confidence(damaged, scan_candidate=True) == 0.3
    result = normalize_document(document_with(page(*native, *ocr, damaged)))
    assert any(d.code == "reconciliation.low_quality_native" for d in result.diagnostics)
    assert "garbled" not in to_markdown(result)


def test_pdf_nested_bullets_preserve_text_and_nonlist_blocks() -> None:
    context = _PdfContext(ParseOptions(), "sample.pdf")
    blocks: list[ContentBlock] = [
        paragraph(str(i), text, box=point_box(x=x, y=i * 15, width=50))
        for i, (text, x) in enumerate((("• Parent", 0), ("• Child", 20), ("• Next", 0)))
    ]
    blocks.insert(0, HeadingBlock(block_id="heading", level=1, spans=(TextSpan(text="Title"),)))
    blocks[1] = cast(ParagraphBlock, blocks[1]).model_copy(
        update={"spans": (TextSpan(text="• ", bold=True), TextSpan(text="Parent"))}
    )
    blocks.append(
        paragraph("tail", "• Footer", box=point_box()).model_copy(
            update={"attributes": {"story": "footer"}}
        )
    )
    output = _paragraphs_and_lists(blocks, context)
    assert isinstance(output[1], ListBlock)
    assert "- Parent\n    - Child\n- Next" in to_markdown(make_document(*output))


def test_pdf_repeated_header_table_keeps_first_occurrence() -> None:
    table = TableBlock(
        block_id="first",
        row_count=1,
        column_count=1,
        source=SourceLocation(page_number=1, bounding_box=point_box(height=5)),
        rows=(
            TableRow(
                row_index=0, cells=(TableCell(column_index=0, displayed_text="Document code"),)
            ),
        ),
    )
    pages = tuple(
        ContainerBlock(
            block_id=f"page{i}",
            role=ContainerRole.PAGE,
            blocks=(table.model_copy(update={"block_id": f"table{i}"}),),
        )
        for i in range(2)
    )
    result = _mark_repeated_tables((*pages, ParagraphBlock(block_id="outside")), 0.6)
    assert to_markdown(make_document(*result)).count("Document code") == 1


def test_pdf_degenerate_table_is_rejected_instead_of_inventing_cells() -> None:
    context = _PdfContext(ParseOptions(), "sample.pdf")
    page_data = _PdfPageData(1, 100, 100, 0, (), (), (), (), False)
    table = _PdfTableData(
        (0, 0, 30, 30),
        (_PdfCellData((0, 0, 0.1, 10), "bad"), _PdfCellData((10, 10, 30, 30), "good")),
    )
    assert _pdf_table(table, page_data, context) is None
    assert context.needs_review


def test_fidelity_preserves_row_span() -> None:
    from document_parser import MarkdownProfile

    table = TableBlock(
        block_id="merged",
        row_count=2,
        column_count=1,
        rows=(
            TableRow(
                row_index=0, cells=(TableCell(column_index=0, row_span=2, displayed_text="Shared"),)
            ),
            TableRow(row_index=1, cells=()),
        ),
    )
    assert 'rowspan="2"' in to_markdown(
        make_document(table), options=MarkdownOptions(profile=MarkdownProfile.FIDELITY)
    )


def test_clean_cell_fallback_and_multiline_formula_never_emit_html() -> None:
    from document_parser import FormulaMode

    table = TableBlock(
        block_id="cell-content",
        row_count=2,
        column_count=1,
        rows=(
            TableRow(
                row_index=0,
                cells=(
                    TableCell(
                        column_index=0,
                        is_header=True,
                        blocks=(
                            ParagraphBlock(block_id="header", spans=(TextSpan(text="Field"),)),
                        ),
                    ),
                ),
            ),
            TableRow(
                row_index=1,
                cells=(TableCell(column_index=0, formula="=SUM(\nA1:A2)", displayed_text="42"),),
            ),
        ),
    )
    document = make_document(table)
    assert "| Field |" in to_markdown(document)
    output = to_markdown(document, options=MarkdownOptions(formula_mode=FormulaMode.FORMULA))
    assert "=SUM(" in output and "A1:A2)" in output and "<br>" not in output


@pytest.mark.parametrize("failure", ("dependency missing", "runtime failed", "wrapped import"))
def test_paddle_initialization_errors_are_typed(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, failure: str
) -> None:
    from document_parser.paddle_ocr import PaddleOcrEngine

    install_modules(monkeypatch, tmp_path, FakePaddleModule({}))

    def broken(_self: object, _module: object) -> object:
        if failure == "wrapped import":
            raise RuntimeError(failure) from ImportError("extra")
        raise RuntimeError(failure)

    monkeypatch.setattr(PaddleOcrEngine, "_build_pipeline", broken)
    error = OcrConfigurationError if failure == "runtime failed" else OcrDependencyNotAvailableError
    with pytest.raises(error):
        PaddleOcrEngine(OcrOptions())
